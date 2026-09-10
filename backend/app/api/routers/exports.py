"""``/exports`` — every data table and analytics view, as a real file.

One endpoint, thirteen entities: ``POST /exports/{entity}`` takes the SAME
filter vocabulary the list endpoints take (``account_group_id``,
``case_file_id``, ``date_from``, ``date_to``) and returns **XLSX** or **PDF**
bytes. PNG is a client-side chart snapshot and is deliberately not served here.

The five rules this router lives by
-----------------------------------
1. **Every matching row, not the current page.** An export that paginated would
   be pointless. The ceiling is :data:`app.services.analytics.EXPORT_ROW_CAP`
   (10 000) and crossing it is a **422** naming the count and the filters to
   narrow — never a silent truncation and never a timeout.
2. **The entity's own View gate.** ``/exports/policies`` needs ``Policies.View``,
   ``/exports/claims`` needs ``Claims.View``. The module is a path parameter, so
   the check is made in the body against ``has_permission`` rather than through a
   statically-bound ``require_permission`` dependency — same matrix, same 403.
3. **Tenant scope, always.** ``broker_id`` is the first predicate of every query
   and every group hop re-asserts it: a foreign ``account_group_id`` exports an
   empty file, it never leaks another broker's rows.
4. **Nothing is filed.** The bytes are streamed straight back. No ``document``
   row, no S3 key in the payload, no S3 key anywhere (non-negotiable 8) — the
   ``document`` table stays the only place an S3 key lives.
5. **One PDF stack.** The branded Playwright pipeline
   (``app.services.pdf`` + ``app.services.pdf_templates.render_table_report_html``)
   renders a landscape table report with the broker's running header, the active
   filters as chips and the row count in the footer.

Latency warning (deploy trap)
-----------------------------
The backend runs behind a **buffered** Lambda/CloudFront invoke, so a very large
PDF render can approach the gateway timeout the same way the comparison align
does. XLSX is cheap; PDF over a few thousand rows is not. Keep the cap, and
prefer XLSX for the wide exports. Do NOT change the invoke mode (OAC constraint 4).
"""
from __future__ import annotations

import io
import re
from datetime import date, datetime, timezone
from decimal import Decimal
from typing import Any

from fastapi import APIRouter, Depends, HTTPException, Response, status
from sqlalchemy import func, select
from sqlalchemy.orm import Session, selectinload

from app.api.deps import get_current_broker_id, get_current_user, get_db
from app.core.case_scope import scope_case_query
from app.core.permissions import has_permission, resolve_role
from app.models.broker import Broker
from app.models.user import User
from app.schemas.analytics import (
    ExportCatalog,
    ExportColumnInfo,
    ExportEntityInfo,
    ExportRequest,
)
from app.services.analytics import (
    ENTITIES,
    ENTITY_KEYS,
    EXPORT_ROW_CAP,
    build_grouped_summary,
    filters_from_payload,
    scope_predicates,
    spec_for,
    validate_group_by,
)

router = APIRouter(prefix="/exports", tags=["exports"])


XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PDF_MIME = "application/pdf"


# ---------------------------------------------------------------------------
# Eager loading — an export walks relationships for every row, so the N+1 is
# real. One selectinload map, kept next to the registry it mirrors.
# ---------------------------------------------------------------------------

def _loader_options(entity: str) -> list[Any]:
    from app.models.case_file import CaseFile
    from app.models.client import Client
    from app.models.collection import CollectionPlan
    from app.models.document import Document
    from app.models.endorsement import Endorsement
    from app.models.inspection import Inspection
    from app.models.offering import Offering
    from app.models.placement import Placement
    from app.models.policy import Claim, Policy
    from app.models.proposal import Proposal
    from app.models.quote import QuoteRequest
    from app.models.sales_lead import SalesLead

    if entity == "case_files":
        return [
            selectinload(CaseFile.client).selectinload(Client.insured),
            selectinload(CaseFile.account_group),
        ]
    if entity == "quotes":
        return [
            selectinload(QuoteRequest.case_file).selectinload(CaseFile.account_group),
            selectinload(QuoteRequest.placement)
            .selectinload(Placement.client)
            .selectinload(Client.insured),
        ]
    if entity == "proposals":
        return [
            selectinload(Proposal.insurer),
            selectinload(Proposal.case_file).selectinload(CaseFile.account_group),
        ]
    if entity == "policies":
        return [
            selectinload(Policy.insurer),
            selectinload(Policy.client).selectinload(Client.insured),
            selectinload(Policy.client).selectinload(Client.account_group),
            selectinload(Policy.case_file),
        ]
    if entity == "endorsements":
        return [
            selectinload(Endorsement.policy),
            selectinload(Endorsement.case_file),
        ]
    if entity == "collections":
        return [
            selectinload(CollectionPlan.policy),
            selectinload(CollectionPlan.case_file),
        ]
    if entity == "claims":
        return [
            selectinload(Claim.policy),
            selectinload(Claim.client).selectinload(Client.insured),
            selectinload(Claim.case_file),
        ]
    if entity == "documents":
        return [selectinload(Document.case_file)]
    if entity == "placements":
        return [
            selectinload(Placement.client).selectinload(Client.insured),
            selectinload(Placement.asset),
            selectinload(Placement.insurance_line),
            selectinload(Placement.case_file),
        ]
    if entity == "inspections":
        return [selectinload(Inspection.asset), selectinload(Inspection.case_file)]
    if entity == "offerings":
        return [selectinload(Offering.quote_request)]
    if entity == "leads":
        return [
            selectinload(SalesLead.insurance_line),
            selectinload(SalesLead.account_group),
        ]
    if entity == "clients":
        return [selectinload(Client.insured), selectinload(Client.account_group)]
    return []


# ---------------------------------------------------------------------------
# Scope + gate
# ---------------------------------------------------------------------------

def _require_view(entity: str, user: User) -> None:
    """Rule 2: each entity's export is gated on that entity's own View grant."""
    spec = spec_for(entity)
    if not has_permission(resolve_role(user), spec.module, "View"):
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail=f"Not permitted to View on {spec.module}",
        )


def _base_predicates(
    db: Session, entity: str, payload: ExportRequest, broker_id: int, user: User
) -> list[Any]:
    """Tenant filter + the scope predicates + the two role-narrowed views.

    ``case_files`` and ``documents`` both honour ``CASE_VIEW_SCOPE``: a process
    desk (``broker_collections`` / ``broker_claims`` / ``broker_inspector``) must
    not be able to export what its own list would hide (v3 rule 8 — one
    visibility mechanism, applied EVERYWHERE, including here).
    """
    from app.api.routers.documents import _case_narrowing_filters
    from app.models.case_file import CaseFile

    spec = spec_for(entity)
    filters = filters_from_payload(payload.filters)
    predicates: list[Any] = [spec.model.broker_id == broker_id]

    if entity == "case_files":
        narrowed = scope_case_query(
            select(CaseFile).where(CaseFile.broker_id == broker_id),
            broker_id=broker_id,
            user=user,
        )
        predicates = [narrowed.whereclause]
    elif entity == "documents":
        predicates.extend(_case_narrowing_filters(db, broker_id, user))

    predicates.extend(scope_predicates(entity, filters, broker_id))

    if payload.filters.status and spec.status_column is not None:
        predicates.append(spec.status_column == payload.filters.status)
    if payload.filters.search and spec.search_predicate is not None:
        predicates.append(spec.search_predicate(payload.filters.search.strip()))
    return predicates


def _resolve_columns(entity: str, requested: list[str] | None):
    spec = spec_for(entity)
    if not requested:
        return list(spec.columns)
    index = {column.key: column for column in spec.columns}
    unknown = [key for key in requested if key not in index]
    if unknown:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail={
                "code": "unknown_columns",
                "message": (
                    f"Columnas desconocidas para '{entity}': {', '.join(unknown)}."
                ),
                "unknown": unknown,
                "available": [column.key for column in spec.columns],
            },
        )
    return [index[key] for key in requested]


# ---------------------------------------------------------------------------
# Value formatting
# ---------------------------------------------------------------------------

def _naive(value: datetime) -> datetime:
    """openpyxl cannot write a tz-aware datetime; normalise to naive UTC."""
    if value.tzinfo is None:
        return value
    return value.astimezone(timezone.utc).replace(tzinfo=None)


def _cell_value(kind: str, value: Any) -> Any:
    """The typed value written into the workbook (dates as dates, UF as floats)."""
    if value is None:
        return None
    if kind == "datetime" and isinstance(value, datetime):
        return _naive(value)
    if kind == "date" and isinstance(value, datetime):
        return value.date()
    if kind in {"uf", "number"}:
        try:
            return float(Decimal(str(value)))
        except Exception:  # noqa: BLE001 - a stray string stays a string
            return str(value)
    if kind == "int":
        try:
            return int(value)
        except Exception:  # noqa: BLE001
            return str(value)
    if kind == "bool":
        return bool(value)
    if isinstance(value, (date, datetime, int, float, Decimal, bool, str)):
        return value
    return str(value)


def _text_value(kind: str, value: Any) -> str:
    """The string written into the PDF table."""
    if value is None:
        return ""
    if kind == "datetime" and isinstance(value, datetime):
        return _naive(value).strftime("%Y-%m-%d %H:%M")
    if kind == "date":
        if isinstance(value, datetime):
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
    if kind == "uf":
        try:
            return f"{Decimal(str(value)):,.4f}".replace(",", "@").replace(".", ",").replace("@", ".")
        except Exception:  # noqa: BLE001
            return str(value)
    if kind == "number":
        try:
            return f"{Decimal(str(value)):,.2f}".replace(",", "@").replace(".", ",").replace("@", ".")
        except Exception:  # noqa: BLE001
            return str(value)
    return str(value)


_NUMERIC_KINDS = {"int", "uf", "number"}

_NUMBER_FORMATS = {
    "date": "yyyy-mm-dd",
    "datetime": "yyyy-mm-dd hh:mm",
    "uf": "#,##0.0000",
    "number": "#,##0.00",
    "int": "0",
}


# ---------------------------------------------------------------------------
# Filename + filter description
# ---------------------------------------------------------------------------

_SAFE = re.compile(r"[^A-Za-z0-9._-]+")


def _scope_slug(payload: ExportRequest) -> str:
    parts: list[str] = []
    if payload.filters.account_group_id:
        parts.append(f"grupo-{payload.filters.account_group_id}")
    if payload.filters.case_file_id:
        parts.append(f"cuenta-{payload.filters.case_file_id}")
    if payload.filters.date_from:
        parts.append(f"desde-{payload.filters.date_from.isoformat()}")
    if payload.filters.date_to:
        parts.append(f"hasta-{payload.filters.date_to.isoformat()}")
    if payload.group_by:
        parts.append(f"por-{payload.group_by}")
    return "-".join(parts) if parts else "completo"


def _filename(entity: str, payload: ExportRequest, extension: str) -> str:
    stem = f"{entity}-{_scope_slug(payload)}-{date.today().isoformat()}"
    return _SAFE.sub("-", stem).strip("-") + f".{extension}"


def _filter_chips(db: Session, entity: str, payload: ExportRequest, broker_id: int):
    """The active filters, resolved to names, for the PDF header."""
    from app.models.account_group import AccountGroup
    from app.models.case_file import CaseFile

    chips: list[dict[str, str]] = []
    f = payload.filters
    if f.account_group_id:
        group = db.scalars(
            select(AccountGroup).where(
                AccountGroup.id == f.account_group_id,
                AccountGroup.broker_id == broker_id,
            )
        ).first()
        chips.append(
            {"label": "Grupo", "value": group.name if group else str(f.account_group_id)}
        )
    if f.case_file_id:
        case = db.scalars(
            select(CaseFile).where(
                CaseFile.id == f.case_file_id, CaseFile.broker_id == broker_id
            )
        ).first()
        chips.append(
            {"label": "Cuenta", "value": case.title if case else str(f.case_file_id)}
        )
    if f.date_from:
        chips.append({"label": "Desde", "value": f.date_from.isoformat()})
    if f.date_to:
        chips.append({"label": "Hasta", "value": f.date_to.isoformat()})
    if f.status:
        chips.append({"label": "Estado", "value": f.status})
    if f.search:
        chips.append({"label": "Búsqueda", "value": f.search})
    if payload.group_by:
        chips.append({"label": "Agrupado por", "value": payload.group_by})
    if not chips:
        chips.append({"label": "Alcance", "value": "Todo el corredor"})
    return chips


# ---------------------------------------------------------------------------
# Row / bucket assembly
# ---------------------------------------------------------------------------

def _count_rows(db: Session, entity: str, predicates: list[Any]) -> int:
    spec = spec_for(entity)
    return int(db.scalar(select(func.count(spec.model.id)).where(*predicates)) or 0)


def _enforce_cap(entity: str, total: int, payload: ExportRequest) -> None:
    if total <= EXPORT_ROW_CAP:
        return
    raise HTTPException(
        status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
        detail={
            "code": "export_too_large",
            "message": (
                f"La exportación de '{entity}' tiene {total} filas y el máximo es "
                f"{EXPORT_ROW_CAP}. Acota los filtros (grupo, grupo-cuenta o fechas)."
            ),
            "entity": entity,
            "row_count": total,
            "row_cap": EXPORT_ROW_CAP,
            "filters": payload.filters.model_dump(mode="json"),
        },
    )


def _fetch_rows(db: Session, entity: str, predicates: list[Any]) -> list[Any]:
    spec = spec_for(entity)
    stmt = (
        select(spec.model)
        .where(*predicates)
        .options(*_loader_options(entity))
        .order_by(spec.order_column.desc())
        .limit(EXPORT_ROW_CAP)
    )
    return list(db.scalars(stmt).unique().all())


# ---------------------------------------------------------------------------
# Writers
# ---------------------------------------------------------------------------

def _write_xlsx(
    *, sheet_title: str, headers: list[tuple[str, str]], rows: list[list[Any]]
) -> bytes:
    """One sheet: header row (frozen, bold), typed cells, autosized columns.

    ``headers`` is ``[(label, kind)]``; ``rows`` carries already-typed values.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = (sheet_title or "Datos")[:31]

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="0E7C70")
    for index, (label, _kind) in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=index, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    widths = [len(label) + 2 for label, _ in headers]
    for row_index, values in enumerate(rows, start=2):
        for column_index, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=column_index, value=value)
            kind = headers[column_index - 1][1]
            number_format = _NUMBER_FORMATS.get(kind)
            if number_format and value is not None:
                cell.number_format = number_format
            if kind in _NUMERIC_KINDS:
                cell.alignment = Alignment(horizontal="right")
            width = len(str(value)) + 2 if value is not None else 3
            if width > widths[column_index - 1]:
                widths[column_index - 1] = width

    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = min(max(width, 8), 48)

    sheet.freeze_panes = "A2"
    if rows:
        sheet.auto_filter.ref = (
            f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
        )

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _branding(db: Session, broker_id: int, running_title: str) -> dict:
    from app.api.routers.antecedentes import _logo_datauri

    broker = db.get(Broker, broker_id)
    return {
        "radal_wordmark": "Radal.",
        "broker_logo_datauri": _logo_datauri(broker) if broker else None,
        "broker_name": (broker.trade_name or broker.legal_name) if broker else None,
        "broker_rut": broker.rut if broker else None,
        "cmf_code": broker.cmf_code if broker else None,
        "cover_meta": [],
        "eyebrow": "Exportación",
        "running_title": running_title,
    }


def _stream(data: bytes, *, mime: str, filename: str, row_count: int) -> Response:
    """Stream the bytes. Never a document row, never an S3 key (rule 4)."""
    return Response(
        content=data,
        media_type=mime,
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Content-Length": str(len(data)),
            "X-Radal-Row-Count": str(row_count),
        },
    )


# ---------------------------------------------------------------------------
# Endpoints
# ---------------------------------------------------------------------------

@router.get("/entities", response_model=ExportCatalog)
def list_export_entities(
    _user: User = Depends(get_current_user),
) -> ExportCatalog:
    """The catalogue the Datos screen builds its column picker from.

    Advertising it keeps the UI honest: a control is either wired or visibly
    disabled (non-negotiable 3), never a button that 404s.
    """
    return ExportCatalog(
        formats=["xlsx", "pdf"],
        row_cap=EXPORT_ROW_CAP,
        entities=[
            ExportEntityInfo(
                entity=spec.entity,
                label=spec.label,
                module=spec.module,
                date_field=spec.date_field,
                money_field=spec.money_field,
                group_by=spec.group_by_keys,
                row_cap=EXPORT_ROW_CAP,
                columns=[
                    ExportColumnInfo(key=c.key, label=c.label, kind=c.kind)
                    for c in spec.columns
                ],
            )
            for spec in ENTITIES.values()
        ],
    )


@router.post("/{entity}")
async def export_entity(
    entity: str,
    payload: ExportRequest,
    db: Session = Depends(get_db),
    broker_id: int = Depends(get_current_broker_id),
    current_user: User = Depends(get_current_user),
) -> Response:
    """Export EVERY row matching the filters as XLSX or PDF.

    ``entity`` is one of ``case_files | quotes | proposals | policies |
    endorsements | collections | claims | documents | placements | inspections |
    offerings | leads | clients``; an unknown one is a 404.

    Body::

        {"format": "xlsx"|"pdf",
         "filters": {"account_group_id": 1, "case_file_id": 2,
                     "date_from": "2026-01-01", "date_to": "2026-12-31",
                     "status": "active", "search": "..."},
         "columns": ["id", "policy_number", ...] | null,
         "group_by": "insurer" | null}

    With ``group_by`` the file carries the aggregated buckets (clave, etiqueta,
    cantidad, total UF) instead of the rows — the same payload the charts read,
    so a PNG snapshot and its XLSX always agree.

    Gated on the entity's own ``View`` grant, scoped to ``broker_id``, capped at
    10 000 rows (422 above it). The response is streamed; nothing is written to
    the ``document`` table.

    **Latency:** the PDF path renders through headless Chromium behind a buffered
    Lambda/CloudFront invoke — a several-thousand-row PDF can run long. XLSX is
    the cheap path for wide exports.
    """
    spec = spec_for(entity)
    _require_view(entity, current_user)
    validate_group_by(entity, payload.group_by)

    predicates = _base_predicates(db, entity, payload, broker_id, current_user)

    if payload.group_by:
        grouped = build_grouped_summary(
            db,
            entity,
            group_by=payload.group_by,
            broker_id=broker_id,
            extra_filters=predicates,
        )
        headers = [("Clave", "text"), ("Etiqueta", "text"), ("Cantidad", "int")]
        if spec.money_field:
            headers.append((f"Total UF ({spec.money_field})", "uf"))
        rows_typed = [
            [b.key, b.label, b.count]
            + ([_cell_value("uf", b.total_uf)] if spec.money_field else [])
            for b in grouped.buckets
        ]
        rows_text = [
            [b.key, b.label, str(b.count)]
            + ([_text_value("uf", b.total_uf)] if spec.money_field else [])
            for b in grouped.buckets
        ]
        row_count = len(grouped.buckets)
        subtitle = (
            f"Agrupado por {payload.group_by} · {grouped.total} registro(s)"
        )
    else:
        total = _count_rows(db, entity, predicates)
        _enforce_cap(entity, total, payload)
        columns = _resolve_columns(entity, payload.columns)
        rows = _fetch_rows(db, entity, predicates)
        headers = [(column.label, column.kind) for column in columns]
        rows_typed = [
            [_cell_value(column.kind, column.get(row)) for column in columns]
            for row in rows
        ]
        rows_text = [
            [_text_value(column.kind, column.get(row)) for column in columns]
            for row in rows
        ]
        row_count = len(rows)
        subtitle = f"{row_count} fila(s)"

    if payload.format == "xlsx":
        data = _write_xlsx(sheet_title=spec.label, headers=headers, rows=rows_typed)
        return _stream(
            data,
            mime=XLSX_MIME,
            filename=_filename(entity, payload, "xlsx"),
            row_count=row_count,
        )

    # --- PDF: the ONE branded Playwright pipeline ---------------------------
    from app.services.pdf import PDFGenerationError, render_html_to_pdf
    from app.services.pdf_templates import render_table_report_html

    html = render_table_report_html(
        title=spec.label,
        branding=_branding(db, broker_id, spec.label),
        columns=[
            {"label": label, "num": kind in _NUMERIC_KINDS} for label, kind in headers
        ],
        rows=rows_text,
        filters=_filter_chips(db, entity, payload, broker_id),
        subtitle=subtitle,
        row_count=row_count,
        footer_note=(
            f"{row_count} fila(s) exportada(s) · "
            f"generado {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M')} UTC"
        ),
    )
    try:
        data = await render_html_to_pdf(html)
    except PDFGenerationError as exc:
        raise HTTPException(
            status_code=status.HTTP_502_BAD_GATEWAY,
            detail=f"No se pudo generar el PDF: {exc}",
        ) from exc

    return _stream(
        data,
        mime=PDF_MIME,
        filename=_filename(entity, payload, "pdf"),
        row_count=row_count,
    )


__all__ = ["router", "ENTITY_KEYS"]
